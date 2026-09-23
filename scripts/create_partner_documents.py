from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "sports-centered"
OUT.mkdir(parents=True, exist_ok=True)

BLUE = "1056D9"
ORANGE = "FF7A00"
DARK = "141820"
MUTED = "5D6675"
LIGHT = "EEF3FB"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def borders(table, color="CCD5E3", size="6"):
    tblPr = table._tbl.tblPr
    el = tblPr.first_child_found_in("w:tblBorders")
    if el is None:
        el = OxmlElement("w:tblBorders")
        tblPr.append(el)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)
        el.append(tag)

def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)

def base_doc(title, subtitle):
    d = Document()
    sec = d.sections[0]
    sec.top_margin = Inches(.52)
    sec.bottom_margin = Inches(.65)
    sec.footer_distance = Inches(.3)
    sec.left_margin = Inches(.72)
    sec.right_margin = Inches(.72)
    styles = d.styles
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(8.7)
    styles["Normal"].font.color.rgb = RGBColor.from_string(DARK)
    styles["Normal"].paragraph_format.space_after = Pt(3)
    styles["Normal"].paragraph_format.line_spacing = 1.0
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor.from_string(DARK)
    styles["Heading 1"].font.size = Pt(13)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor.from_string(BLUE)
    styles["Heading 1"].paragraph_format.space_before = Pt(6)
    styles["Heading 1"].paragraph_format.space_after = Pt(2)
    styles["Heading 2"].font.size = Pt(10.5)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].font.color.rgb = RGBColor.from_string(DARK)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("KOBE’S BETTING HUB")
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string(ORANGE)
    p = d.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(MUTED)
    return d

def clause(d, n, heading, body):
    p = d.add_paragraph(style="Heading 1")
    p.add_run(f"{n}. {heading}")
    d.add_paragraph(body)

def add_bullets(d, items):
    for item in items:
        p = d.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(.25)
        p.paragraph_format.first_line_indent = Inches(-.13)
        p.add_run(item)

def build_agreement():
    d = base_doc("CREATOR & AFFILIATE PARTNERSHIP AGREEMENT", "Draft for signature • Sports Centered")
    t = d.add_table(rows=4, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Inches(1.75); t.columns[1].width = Inches(5.25)
    rows = [
        ("Effective date", "____________________"),
        ("KBH", "Kobe’s Betting Hub, operated by ______________________________"),
        ("Partner", "Sports Centered, operated by ______________________________"),
        ("Partner contact", "Name: ____________________  Email: ____________________"),
    ]
    for row, (a, b) in zip(t.rows, rows):
        row.cells[0].text = a; row.cells[1].text = b
        shade(row.cells[0], LIGHT)
        for c in row.cells: set_cell_margin(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(t)
    d.add_paragraph()
    d.add_paragraph("This Agreement states the complete terms for Partner’s promotion of Kobe’s Betting Hub (“KBH”). It is intended as a practical business draft and should be reviewed by qualified counsel before signature.")

    clause(d, 1, "Appointment and scope", "KBH appoints Partner on a non-exclusive, revocable basis to promote KBH through Partner’s owned or authorized sports-media channels. Partner may create original variations, but no advertisement, caption, landing-page claim, testimonial, or material edit may go live until KBH approves the final version in writing.")
    clause(d, 2, "Approved tracking", "KBH will issue Partner a unique referral link and/or code. Partner must use that exact link or code in every promotion. Attribution applies only when a new customer completes a qualifying paid VIP membership through the assigned tracking path and the code is attached to checkout. An entered valid Partner code may override other campaign attribution. Lost, removed, altered, blocked, cross-device, or untracked journeys do not create a commission unless KBH can reasonably verify attribution from its records.")
    clause(d, 3, "Commission", "KBH will pay Partner a one-time $10.00 commission for each new qualifying paid VIP subscriber attributed to Partner; renewals do not earn additional commission unless the parties later agree in writing. A conversion qualifies only after KBH receives and retains the subscriber’s first successful paid membership charge through a commission-eligible offer and it passes a seven-day refund, dispute, and fraud review. Free accounts, trials that never convert, duplicate accounts, self-referrals, existing or returning customers, test transactions, complimentary access, fraud, refunds, disputes, and chargebacks do not qualify. A single customer/payment may earn only one partner commission.")
    clause(d, 4, "Reconciliation and payment", "KBH will reconcile qualifying conversions monthly and provide a partner-specific statement. Earned commissions will be paid after the applicable refund/chargeback review period and once any required payout and tax information is complete. Reversed or later-disputed payments may be withheld, offset against future commissions, or recovered where permitted by law. Partner must raise a statement dispute within 30 days after delivery.")
    clause(d, 5, "Creator and VIP access", "During the active partnership, KBH may provide designated Partner personnel complimentary creator/VIP Discord access solely to review the service and create accurate promotions. This access is not a paid membership, is non-transferable, creates no commission, and remains active only while this Agreement is in effect. KBH may suspend or remove it immediately for misuse, security concerns, inaccurate promotion, or termination. Partner will promptly identify anyone who no longer needs access.")
    clause(d, 6, "Advertising standards and disclosures", "Partner will make truthful, supportable statements and clearly disclose the paid relationship close to each endorsement (for example, “Paid partnership with Kobe’s Betting Hub” or “Ad — I may earn a commission for paid memberships through this link”). Partner will use available platform disclosure tools in addition to, not instead of, a clear disclosure in the content. Partner will not claim or imply guaranteed winnings, risk-free betting, certain profit, insider information, fixed games, or a specific win rate unless KBH has supplied current written substantiation and approved the exact claim.")
    clause(d, 7, "Audience and responsible-gaming rules", "Partner will not knowingly target minors or persons below the legal betting age, will not target places where the promotion or service is unlawful, and will follow applicable platform, advertising, sports-wagering, and responsible-gaming requirements. Each promotion must include “21+ • Where legal • Bet responsibly • No guaranteed outcomes,” or another written form approved by KBH. Partner will not encourage chasing losses, borrowing to bet, or betting beyond a person’s means.")
    clause(d, 8, "Brand and content license", "KBH grants Partner a limited, non-exclusive, non-transferable, revocable license during the Term to use approved KBH names, logos, and assets only for this partnership. Partner may not alter the logo, register confusingly similar names, buy search terms using KBH’s marks without written approval, or imply ownership or agency. Partner retains its original content; Partner grants KBH a non-exclusive, worldwide, royalty-free license during the Term and for 12 months afterward to repost approved partnership content with attribution.")
    clause(d, 9, "Records, reporting, and data", "KBH’s tracking and payment records control absent clear error. KBH will provide Partner only Partner-specific activity and commission information; Partner receives no access to KBH’s full administrative dashboard, customer personal data, or other partners’ results. Each party will protect confidential information and use personal data only as authorized and legally permitted.")
    clause(d, 10, "Independent contractor; taxes", "Partner is an independent contractor and has no authority to bind KBH. Partner is responsible for its personnel, expenses, insurance, licenses, and taxes. KBH may require a completed Form W-9 and may issue tax reporting required by law.")
    clause(d, 11, "Term and termination", "This Agreement begins on the Effective Date and continues month-to-month. Either party may terminate it by written notice. KBH may terminate or suspend immediately for fraud, unauthorized claims, missing disclosures, legal or platform risk, brand misuse, data or security concerns, or material breach. On termination, Partner will stop using links and brand assets, remove scheduled promotions, and lose complimentary creator/VIP access. Valid commissions earned before termination remain payable subject to this Agreement’s exclusions and review period.")
    clause(d, 12, "Risk allocation", "Neither party promises any volume, conversion rate, audience response, revenue, betting outcome, or continued program availability. To the maximum extent permitted by law, neither party is liable for indirect, special, or consequential damages. Each party will be responsible for claims arising from its own breach, unlawful conduct, or unauthorized statements. Any broader indemnity, liability cap, or insurance requirement should be completed with counsel before signature.")
    clause(d, 13, "General terms", "This Agreement and approved written campaign instructions are the entire agreement on this partnership and may be changed only in writing accepted by both parties. Neither party may assign it without the other’s written consent, except in a business reorganization. If a provision is unenforceable, the remainder survives. Notices may be sent to the contact emails above. Governing law and venue: State of ____________________, County of ____________________. Electronic signatures and counterparts are permitted.")

    d.add_heading("Campaign schedule", level=1)
    s = d.add_table(rows=6, cols=2)
    s.alignment = WD_TABLE_ALIGNMENT.CENTER
    data = [
        ("Commission", "$10.00 per qualifying new paid VIP subscriber"),
        ("Current consumer offer", "$19.99/month through October 22; then $32.99/month, subject to KBH’s written updates"),
        ("Tracking link/code", "To be issued by KBH; use the exact assigned link/code"),
        ("Reconciliation", "Monthly; partner-specific statement"),
        ("Creative approval", "Written approval required before first publication and material edits"),
        ("Complimentary access", "Active partnership only; removable at termination or for cause"),
    ]
    for row, pair in zip(s.rows, data):
        row.cells[0].text, row.cells[1].text = pair
        shade(row.cells[0], LIGHT)
        for c in row.cells: set_cell_margin(c)
    borders(s)

    d.add_heading("Signatures", level=1)
    sig = d.add_table(rows=4, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_data = [
        ("KOBE’S BETTING HUB", "SPORTS CENTERED"),
        ("By: ______________________________", "By: ______________________________"),
        ("Name/Title: _______________________", "Name/Title: _______________________"),
        ("Date: _____________________________", "Date: _____________________________"),
    ]
    for row, pair in zip(sig.rows, sig_data):
        row.cells[0].text, row.cells[1].text = pair
        for c in row.cells: set_cell_margin(c, 120, 140, 120, 140)
    borders(sig, color="FFFFFF", size="0")
    footer = d.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Kobe’s Betting Hub • Creator & Affiliate Partnership Agreement")
    r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(MUTED)
    d.save(OUT / "Sports_Centered_Affiliate_Agreement.docx")

def build_kit():
    d = base_doc("SPORTS CENTERED PROMOTION KIT", "Approved starting copy • Final posts still require KBH written approval")
    d.add_heading("The offer", level=1)
    add_bullets(d, [
        "Daily sports picks, full pick writeups, tracked results, free Discord access, and paid VIP membership.",
        "$19.99/month through October 22; regular monthly price $32.99 afterward.",
        "Use only the exact Sports Centered tracking link/code supplied by KBH.",
    ])
    d.add_heading("Approved primary caption", level=1)
    p = d.add_paragraph()
    p.paragraph_format.left_indent = Inches(.25)
    p.paragraph_format.right_indent = Inches(.25)
    p.add_run("Paid partnership with Kobe’s Betting Hub.\n\nDaily picks, full writeups, and transparent tracked results—plus free Discord access and an optional VIP membership. VIP is $19.99/month through October 22, then $32.99/month.\n\nJoin through our link: [SPORTS CENTERED TRACKING LINK]\n\n21+ • Where legal • Bet responsibly • No guaranteed outcomes").bold = True
    d.add_heading("Short story / post copy", level=1)
    d.add_paragraph("Ad — Daily picks. Full writeups. Tracked results. VIP $19.99/month through October 22. Join Kobe’s Betting Hub through our link. 21+ • Where legal • Bet responsibly • No guaranteed outcomes.")
    d.add_heading("Creative guardrails", level=1)
    add_bullets(d, [
        "Send the final image/video and exact caption to KBH for written approval before it goes live.",
        "Keep the paid-partnership disclosure visible and close to the endorsement; do not bury it after “more.”",
        "Do not say guaranteed, lock, risk-free, can’t lose, fixed, insider, automatic profit, or promise a win rate.",
        "Do not alter the price, deadline, tracking link, logo, responsible-gaming footer, or membership benefits without KBH approval.",
        "If a platform offers a Paid Partnership label, use it and keep the written disclosure too.",
    ])
    d.add_heading("Tracking and reporting", level=1)
    d.add_paragraph("Sports Centered receives its own link/code. KBH tracks partner-attributed visits, checkout starts, successful paid VIP memberships, reversals, and commission status. KBH will send a partner-only monthly statement. Sports Centered should not receive access to KBH’s complete admin dashboard, customer data, or other partners’ performance.")
    d.add_heading("Approval reply format", level=1)
    d.add_paragraph("Send: (1) final creative, (2) exact caption, (3) platform and placement, (4) intended date/time, and (5) the exact tracking link. KBH replies APPROVED or lists required changes. Any material edit requires reapproval.")
    d.save(OUT / "Sports_Centered_Promotion_Kit.docx")

if __name__ == "__main__":
    build_agreement()
    build_kit()
